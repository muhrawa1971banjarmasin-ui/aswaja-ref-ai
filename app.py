import re
from io import BytesIO

import streamlit as st
from docx import Document
from pypdf import PdfReader

st.set_page_config(page_title="ASWAJA REF-AI", page_icon="📚", layout="wide")

st.title("📚 ASWAJA REF-AI")
st.caption("Asisten Akademik Berbasis Generative AI dengan Sistem Validasi Referensi Aswaja An-Nahdliyah dan Sitasi Otomatis Karya Ilmiah Mahasiswa")
st.info("Versi 0.1: pembacaan dokumen dan deteksi awal. Validasi sumber asli, Generative AI, dan sitasi otomatis masih dalam pengembangan.")

with st.sidebar:
    st.header("Menu Aplikasi")
    menu = st.radio("Pilih menu", ["Beranda", "Unggah Karya Ilmiah", "Validasi Referensi Aswaja", "Sitasi Otomatis", "Laporan Pengembangan"])

if "isi_dokumen" not in st.session_state:
    st.session_state.isi_dokumen = ""
if "nama_dokumen" not in st.session_state:
    st.session_state.nama_dokumen = ""


def baca_dokumen(file):
    data = BytesIO(file.getvalue())
    if file.name.lower().endswith(".docx"):
        dok = Document(data)
        teks_paragraf = [p.text for p in dok.paragraphs]
        teks_tabel = [cell.text for tabel in dok.tables for baris in tabel.rows for cell in baris.cells]
        return "\n".join(teks_paragraf + teks_tabel)
    pdf = PdfReader(data)
    return "\n".join(halaman.extract_text() or "" for halaman in pdf.pages)


def temukan_kutipan(teks):
    pola = r"\([^()\n]{2,90}?,\s*(?:19|20)\d{2}(?:\s*[:,]\s*[^()\n]{1,25})?\)"
    hasil = re.findall(pola, teks)
    return list(dict.fromkeys(hasil))


if menu == "Beranda":
    st.subheader("Selamat datang")
    st.write("Aplikasi ini dikembangkan sebagai produk penelitian R&D dengan model ADDIE. Fokusnya adalah validasi literatur dan referensi Aswaja An-Nahdliyah serta sitasi otomatis.")
    st.write("**Alur yang direncanakan:** Unggah dokumen → identifikasi kutipan → verifikasi sumber Aswaja → rekomendasi → sitasi otomatis → laporan.")
    st.warning("Hasil deteksi awal bukan bukti bahwa sebuah sumber valid atau tidak valid.")

elif menu == "Unggah Karya Ilmiah":
    st.subheader("📤 Unggah Karya Ilmiah")
    unggahan = st.file_uploader("Pilih satu dokumen PDF atau DOCX", type=["pdf", "docx"])
    if unggahan is not None:
        try:
            isi = baca_dokumen(unggahan)
            if not isi.strip():
                st.warning("Teks tidak terbaca. PDF hasil pemindaian mungkin memerlukan OCR pada versi berikutnya.")
            else:
                st.session_state.isi_dokumen = isi
                st.session_state.nama_dokumen = unggahan.name
                st.success("Dokumen berhasil dibaca.")
                st.write("**Nama dokumen:**", unggahan.name)
                st.write("**Jumlah kata (perkiraan):**", len(isi.split()))
                kutipan = temukan_kutipan(isi)
                st.write("**Kutipan pola penulis-tahun yang terdeteksi:**", len(kutipan))
                if kutipan:
                    st.write("**Contoh hasil deteksi awal:**")
                    for nomor, kutip in enumerate(kutipan[:30], 1):
                        st.write(f"{nomor}. {kutip}")
                else:
                    st.info("Belum ada kutipan pola penulis-tahun yang terdeteksi. Catatan kaki dan gaya kutipan lain belum didukung penuh.")
                with st.expander("Lihat cuplikan teks dokumen"):
                    st.text_area("Cuplikan", isi[:5000], height=240, disabled=True)
        except Exception as exc:
            st.error(f"Dokumen belum dapat dibaca: {exc}")

elif menu == "Validasi Referensi Aswaja":
    st.subheader("🔎 Validasi Referensi Aswaja An-Nahdliyah")
    if st.session_state.isi_dokumen:
        st.write("Dokumen aktif:", st.session_state.nama_dokumen)
        st.write("Kutipan pola penulis-tahun yang terdeteksi:", len(temukan_kutipan(st.session_state.isi_dokumen)))
    else:
        st.info("Unggah karya ilmiah terlebih dahulu melalui menu Unggah Karya Ilmiah.")
    st.warning("Belum tersedia pemeriksaan sumber asli. Aplikasi tidak menyatakan kutipan valid, tidak valid, atau palsu pada versi ini.")
    st.write("Tahap selanjutnya: bank referensi Aswaja, pencarian bibliografi, pencocokan teks sumber, dan analisis berbasis Generative AI.")

elif menu == "Sitasi Otomatis":
    st.subheader("📝 Sitasi Otomatis")
    gaya = st.selectbox("Gaya sitasi yang direncanakan", ["Chicago (Catatan Kaki dan Bibliografi)", "APA Edisi 7"])
    st.write("Gaya dipilih:", gaya)
    st.info("Pembuatan sitasi otomatis akan diaktifkan setelah data sumber dapat diverifikasi. Sistem tidak akan mengarang penulis, tahun, halaman, DOI, atau judul.")

elif menu == "Laporan Pengembangan":
    st.subheader("📊 Laporan Pengembangan R&D")
    st.write("**Model pengembangan:** ADDIE")
    st.write("**Versi aplikasi:** 0.1 (prototipe awal)")
    st.write("**Fitur yang disiapkan:** antarmuka, unggah PDF/DOCX, ekstraksi teks, deteksi awal pola kutipan penulis-tahun.")
    st.write("**Fitur yang belum selesai:** validasi sumber, analisis AI, catatan kaki, rekomendasi sumber, sitasi otomatis, ekspor laporan.")
    st.caption("Keberhasilan fitur baru dapat dicatat setelah pengujian nyata.")

st.divider()
st.caption("ASWAJA REF-AI | Prototipe penelitian R&D | Tidak menggantikan verifikasi akademik oleh ahli")
